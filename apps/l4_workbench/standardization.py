"""Turn immutable source snapshots into traceable document and question assets."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import mimetypes
import re
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4
from xml.etree import ElementTree as ET

from .domain import ValidationError
from .manifest_adapter import LocalManifestAdapter
from .pipeline import sha256_file
from .tagging import attach_unit_tag_profiles, build_tag_profile


PARSER_VERSION = "standardizer-v2"
QUESTION_START = re.compile(r"^\s*(?:第\s*)?(\d{1,3})\s*[.、．)）]\s*")
UNIT_PATTERN = re.compile(r"(?:（([一二三四五六七八九十百\d]+)）|\(([一二三四五六七八九十百\d]+)\))")
IMAGE_HINT = re.compile(r"\[(?:图片|图|image)[:：]?", re.IGNORECASE)
ANSWER_SECTION_HEADING = re.compile(r"(?m)^\s*(参考答案与解析|答案与解析|答案及解析|答案解析)\s*[:：]?\s*(?:$|\n)")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def stable_id(prefix: str, *parts: Any) -> str:
    digest = hashlib.sha256("\u241f".join(str(part) for part in parts).encode("utf-8")).hexdigest()
    return f"{prefix}-{digest[:14]}"


def normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value).lower()


def apply_answer_section_boundary(blocks: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    """Exclude an answer section only when its heading begins a standalone line."""
    for block in blocks:
        if block.get("type") not in {"text", "formula"}:
            continue
        match = ANSWER_SECTION_HEADING.search(str(block.get("text") or ""))
        if not match:
            continue
        metadata = {
            "heading": match.group(1),
            "source_locator": block.get("source_locator"),
            "page": block.get("page"),
        }
        if block.get("page") is not None:
            cutoff_page = int(block["page"])
            return [item for item in blocks if int(item.get("page", cutoff_page)) < cutoff_page], metadata
        kept = [item for item in blocks if int(item.get("sequence", 0)) < int(block.get("sequence", 0))]
        prefix = str(block.get("text") or "")[:match.start()].strip()
        if prefix:
            partial = dict(block)
            partial["text"] = prefix
            kept.append(partial)
        return kept, metadata
    return blocks, None


class DocumentStandardizer:
    """Parse local snapshot files while retaining source and visual evidence."""

    def __init__(self, output_root: Path) -> None:
        self.output_root = Path(output_root).resolve()
        self.derived_root = self.output_root / "derived"

    def standardize(self, snapshot: dict[str, Any]) -> dict[str, Any]:
        if snapshot.get("status") not in {"frozen", "partial"}:
            raise ValidationError("only frozen source snapshots can be standardized")
        run_id = stable_id("standardization", snapshot["immutable_checksum"], PARSER_VERSION)
        documents: list[dict[str, Any]] = []
        question_assets: list[dict[str, Any]] = []
        issues: list[dict[str, Any]] = []
        run_root = self.derived_root / run_id
        run_root.mkdir(parents=True, exist_ok=True)

        for file_record in snapshot.get("files", []):
            source = (self.output_root / file_record["relative_path"]).resolve()
            if not source.is_relative_to(self.output_root) or not source.is_file():
                issues.append(self._issue("error", "source_file_missing", f"快照文件不存在：{file_record.get('name')}", None))
                continue
            document, document_issues = self._parse_file(snapshot, file_record, source, run_root)
            documents.append(document)
            issues.extend(document_issues)
            assets, asset_issues = self._segment_document(snapshot, document)
            question_assets.extend(assets)
            issues.extend(asset_issues)

        status = "completed"
        if any(issue["severity"] == "error" for issue in issues):
            status = "completed_with_blockers"
        elif issues:
            status = "completed_with_issues"
        payload = {
            "id": run_id,
            "source_snapshot_id": snapshot["id"],
            "source_checksum": snapshot["immutable_checksum"],
            "parser_version": PARSER_VERSION,
            "status": status,
            "document_ids": [item["id"] for item in documents],
            "question_asset_ids": [item["id"] for item in question_assets],
            "document_count": len(documents),
            "question_asset_count": len(question_assets),
            "issue_count": len(issues),
            "issues": issues,
            "created_at": now(),
        }
        payload["result_checksum"] = hashlib.sha256(
            json.dumps({"documents": documents, "question_assets": question_assets, "issues": issues}, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()
        return {"run": payload, "documents": documents, "question_assets": question_assets}

    def standardize_base_records(
        self,
        snapshot: dict[str, Any],
        base_export: dict[str, Any],
        mapping: dict[str, str],
    ) -> dict[str, Any]:
        run_id = stable_id("standardization", snapshot["immutable_checksum"], PARSER_VERSION, json.dumps(mapping, sort_keys=True, ensure_ascii=False))
        document_id = stable_id("document", snapshot["id"], base_export.get("base_token"), base_export.get("table_id"), base_export.get("view_id"))
        document = {
            "id": document_id,
            "source_snapshot_id": snapshot["id"],
            "source_kind": "feishu_base",
            "source_name": base_export.get("table_name") or base_export.get("table_id"),
            "source_file_sha256": snapshot["immutable_checksum"],
            "record_count": len(base_export.get("records", [])),
            "content_blocks": [],
            "parser_version": PARSER_VERSION,
        }
        assets = []
        issues = []
        for index, record in enumerate(base_export.get("records", []), 1):
            raw_fields = record.get("fields", {})
            text = self._mapped_value(raw_fields, mapping.get("question_text"))
            text = self._plain_value(text)
            image_value = self._mapped_value(raw_fields, mapping.get("question_image"))
            blocks = []
            if text:
                blocks.append({"type": "text", "text": text, "sequence": 1, "source_locator": f"record:{record['record_id']}"})
            if image_value not in (None, "", []):
                blocks.append({
                    "type": "image_reference", "value": image_value, "sequence": len(blocks) + 1,
                    "source_locator": f"record:{record['record_id']}", "materialized": False,
                })
            asset = self._asset_from_blocks(
                snapshot, document, blocks, index, str(record.get("record_id")),
                source_fields=raw_fields,
                normalized_fields={canonical: self._mapped_value(raw_fields, source) for canonical, source in mapping.items() if source},
            )
            if not text:
                issue = self._issue("error", "question_text_missing", "记录缺少可映射的题目文本", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
            if image_value not in (None, "", []):
                issue = self._issue("error", "remote_image_not_materialized", "题目截图仍是远端引用，正式生产前必须下载并校验", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
                asset["image_integrity"] = "remote_reference_unmaterialized"
            elif IMAGE_HINT.search(text or ""):
                issue = self._issue("error", "declared_image_missing", "题干声明含图，但未找到题目截图字段", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
                asset["image_integrity"] = "missing"
            assets.append(asset)

        status = "completed_with_blockers" if any(x["severity"] == "error" for x in issues) else "completed"
        run = {
            "id": run_id, "source_snapshot_id": snapshot["id"], "source_checksum": snapshot["immutable_checksum"],
            "parser_version": PARSER_VERSION, "status": status, "document_ids": [document_id],
            "question_asset_ids": [item["id"] for item in assets], "document_count": 1,
            "question_asset_count": len(assets), "issue_count": len(issues), "issues": issues, "created_at": now(),
        }
        run["result_checksum"] = hashlib.sha256(json.dumps({"assets": assets, "issues": issues}, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        return {"run": run, "documents": [document], "question_assets": assets}

    def standardize_manifest_records(
        self,
        snapshot: dict[str, Any],
        manifest_path: Path,
        records: list[dict[str, Any]],
        mapping: dict[str, str],
    ) -> dict[str, Any]:
        """Build question assets from a local row manifest and its frozen images."""
        mapping_json = json.dumps(mapping, sort_keys=True, ensure_ascii=False)
        run_id = stable_id("standardization", snapshot["immutable_checksum"], PARSER_VERSION, mapping_json)
        document_id = stable_id("document", snapshot["id"], manifest_path.name)
        document = {
            "id": document_id,
            "source_snapshot_id": snapshot["id"],
            "source_kind": "local_manifest",
            "source_name": manifest_path.name,
            "source_file_sha256": snapshot["immutable_checksum"],
            "record_count": len(records),
            "content_blocks": [],
            "parser_version": PARSER_VERSION,
        }
        captured_by_sha = {
            item["sha256"]: item
            for item in snapshot.get("files", [])
            if item.get("extension") in {".png", ".jpg", ".jpeg", ".webp"}
        }
        adapter = LocalManifestAdapter()
        assets = []
        issues = []
        image_field = mapping.get("question_image")
        for index, record in enumerate(records, 1):
            record_id = self._plain_value(record.get(mapping.get("source_id"))) or adapter.record_id(record, index)
            text = self._plain_value(record.get(mapping.get("question_text")))
            blocks = []
            if text:
                blocks.append({"type": "text", "text": text, "sequence": 1, "source_locator": f"record:{record_id}"})
            image_values = adapter.path_values(record.get(image_field)) if image_field else []
            missing_image_values = []
            suspicious_images = []
            for raw_value in image_values:
                source_image = adapter.resolve_image_path(raw_value, manifest_path.parent)
                captured = None
                if source_image and source_image.is_file():
                    captured = captured_by_sha.get(sha256_file(source_image))
                if captured:
                    blocks.append({
                        "type": "image", "path": captured["relative_path"], "sha256": captured["sha256"],
                        "mime_type": mimetypes.guess_type(captured["name"])[0] or "application/octet-stream",
                        "sequence": len(blocks) + 1, "source_locator": f"record:{record_id}",
                    })
                    inspection = adapter.inspect_image(source_image)
                    if inspection.get("suspicious"):
                        suspicious_images.append(inspection)
                else:
                    missing_image_values.append(raw_value)
                    blocks.append({
                        "type": "image_reference", "value": raw_value, "materialized": False,
                        "sequence": len(blocks) + 1, "source_locator": f"record:{record_id}",
                    })
            asset = self._asset_from_blocks(
                snapshot, document, blocks, index, str(record_id), source_fields=record,
                normalized_fields={canonical: record.get(source) for canonical, source in mapping.items() if source},
            )
            mapped_paper = self._plain_value(record.get(mapping.get("source_paper")))
            mapped_question_no = self._plain_value(record.get(mapping.get("question_no")))
            if mapped_paper:
                asset["source_name"] = mapped_paper
            if mapped_question_no:
                asset["question_no"] = mapped_question_no
            if not text:
                issue = self._issue("error", "question_text_missing", "记录缺少可映射的题目文本", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
            if missing_image_values:
                fully_missing = not any(block["type"] == "image" for block in blocks)
                code = "local_image_missing" if fully_missing else "local_image_partial"
                message = "清单引用的本地题图不存在或未进入冻结快照" if fully_missing else "清单中的部分本地题图缺失"
                issue = self._issue("error", code, message, document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
                asset["image_integrity"] = "missing" if fully_missing else "partial"
            elif not image_values and IMAGE_HINT.search(text or ""):
                issue = self._issue("error", "declared_image_missing", "题干声明含图，但清单没有可用题图字段", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
                asset["image_integrity"] = "missing"
            if suspicious_images:
                issue = self._issue("warning", "image_extreme_aspect_ratio", "题图纵横比异常，可能误裁为整份试卷", document_id, asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
            assets.append(asset)

        status = "completed_with_blockers" if any(item["severity"] == "error" for item in issues) else ("completed_with_issues" if issues else "completed")
        run = {
            "id": run_id, "source_snapshot_id": snapshot["id"], "source_checksum": snapshot["immutable_checksum"],
            "parser_version": PARSER_VERSION, "status": status, "document_ids": [document_id],
            "question_asset_ids": [item["id"] for item in assets], "document_count": 1,
            "question_asset_count": len(assets), "issue_count": len(issues), "issues": issues, "created_at": now(),
        }
        run["result_checksum"] = hashlib.sha256(
            json.dumps({"assets": assets, "issues": issues}, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()
        return {"run": run, "documents": [document], "question_assets": assets}

    def _parse_file(
        self,
        snapshot: dict[str, Any],
        file_record: dict[str, Any],
        source: Path,
        run_root: Path,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        document_id = stable_id("document", snapshot["id"], file_record["sha256"])
        target_root = run_root / document_id
        target_root.mkdir(parents=True, exist_ok=True)
        extension = source.suffix.lower()
        issues: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []
        if extension == ".docx":
            blocks, issues = self._parse_docx(source, target_root, document_id)
        elif extension == ".pdf":
            blocks, issues = self._parse_pdf(source, target_root, document_id)
        elif extension in {".png", ".jpg", ".jpeg", ".webp"}:
            blocks = [{
                "type": "image", "path": file_record["relative_path"], "sha256": file_record["sha256"],
                "mime_type": mimetypes.guess_type(source.name)[0] or "application/octet-stream", "sequence": 1,
                "source_locator": "whole-file",
            }]
        elif extension in {".txt", ".md"}:
            blocks = self._text_blocks(source.read_text(encoding="utf-8", errors="replace"))
        elif extension == ".csv":
            with source.open(encoding="utf-8-sig", errors="replace", newline="") as stream:
                rows = list(csv.reader(stream))
            blocks = [{"type": "table", "rows": rows, "sequence": 1, "source_locator": "sheet:1"}]
        elif extension == ".json":
            value = json.loads(source.read_text(encoding="utf-8"))
            blocks = [{"type": "text", "text": json.dumps(value, ensure_ascii=False, indent=2), "sequence": 1, "source_locator": "json-root"}]
        elif extension == ".xlsx":
            blocks, issues = self._parse_xlsx(source, document_id)
        elif extension in {".doc", ".xls"}:
            issues = [self._issue("error", "legacy_format_needs_conversion", f"{extension} 需先转换为 OOXML 格式", document_id)]
        else:
            issues = [self._issue("error", "unsupported_parser", f"尚未实现解析：{extension}", document_id)]

        blocks, answer_boundary = apply_answer_section_boundary(blocks)
        document = {
            "id": document_id, "source_snapshot_id": snapshot["id"], "source_kind": "local_file",
            "source_name": file_record["name"], "source_path": file_record["relative_path"],
            "source_file_sha256": file_record["sha256"], "mime_type": mimetypes.guess_type(source.name)[0],
            "content_blocks": blocks, "parser_version": PARSER_VERSION,
            "answer_section_boundary": answer_boundary,
        }
        return document, issues

    def _parse_docx(self, source: Path, target_root: Path, document_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        issues: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []
        with zipfile.ZipFile(source) as archive:
            try:
                root = ET.fromstring(archive.read("word/document.xml"))
            except (KeyError, ET.ParseError) as error:
                return [], [self._issue("error", "docx_document_xml_invalid", str(error), document_id)]
            relationships: dict[str, str] = {}
            try:
                rel_root = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
                relationships = {item.attrib["Id"]: item.attrib.get("Target", "") for item in rel_root.findall(f"{{{REL}}}Relationship")}
            except (KeyError, ET.ParseError):
                pass
            body = root.find(f"{{{W}}}body")
            sequence = 0
            for element in list(body) if body is not None else []:
                if element.tag == f"{{{W}}}p":
                    texts = [node.text or "" for node in element.iter() if node.tag in {f"{{{W}}}t", f"{{{M}}}t"}]
                    text = "".join(texts).strip()
                    if text:
                        sequence += 1
                        block_type = "formula" if element.find(f".//{{{M}}}oMath") is not None else "text"
                        blocks.append({"type": block_type, "text": text, "sequence": sequence, "source_locator": f"paragraph:{sequence}"})
                    for blip in element.findall(f".//{{{A}}}blip"):
                        relation_id = blip.attrib.get(f"{{{R}}}embed")
                        target = relationships.get(relation_id or "", "")
                        archive_name = str(Path("word") / target).replace("\\", "/")
                        if not target or archive_name not in archive.namelist():
                            issues.append(self._issue("error", "docx_image_relationship_missing", f"无法解析图片关系：{relation_id}", document_id))
                            continue
                        media_name = Path(target).name
                        media_path = target_root / "media" / media_name
                        media_path.parent.mkdir(parents=True, exist_ok=True)
                        media_path.write_bytes(archive.read(archive_name))
                        sequence += 1
                        blocks.append({
                            "type": "image", "path": self._relative(media_path), "sha256": sha256_file(media_path),
                            "mime_type": mimetypes.guess_type(media_name)[0] or "application/octet-stream",
                            "sequence": sequence, "source_locator": f"paragraph-image:{sequence}",
                        })
                elif element.tag == f"{{{W}}}tbl":
                    rows = []
                    for row in element.findall(f"{{{W}}}tr"):
                        cells = []
                        for cell in row.findall(f"{{{W}}}tc"):
                            cells.append("".join(node.text or "" for node in cell.iter(f"{{{W}}}t")).strip())
                        rows.append(cells)
                    sequence += 1
                    blocks.append({"type": "table", "rows": rows, "sequence": sequence, "source_locator": f"table:{sequence}"})
        return blocks, issues

    def _parse_pdf(self, source: Path, target_root: Path, document_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        issues: list[dict[str, Any]] = []
        page_texts: list[str] = []
        if importlib.util.find_spec("pypdf"):
            try:
                from pypdf import PdfReader
                page_texts = [(page.extract_text() or "").strip() for page in PdfReader(str(source)).pages]
            except Exception as error:
                issues.append(self._issue("warning", "pdf_text_extraction_failed", str(error), document_id))
        else:
            issues.append(self._issue("warning", "pdf_text_dependency_missing", "缺少 pypdf；已保留页面图，但文本需补 OCR", document_id))

        page_images: list[Path] = []
        renderer = shutil.which("pdftoppm")
        if renderer:
            prefix = target_root / "pages" / "page"
            prefix.parent.mkdir(parents=True, exist_ok=True)
            try:
                subprocess.run([renderer, "-png", "-r", "120", str(source), str(prefix)], check=True, timeout=180, capture_output=True)
                page_images = sorted(prefix.parent.glob("page-*.png"), key=self._page_number)
            except (subprocess.SubprocessError, OSError) as error:
                issues.append(self._issue("error", "pdf_page_render_failed", str(error), document_id))
        else:
            issues.append(self._issue("error", "pdf_renderer_missing", "缺少 pdftoppm，无法冻结 PDF 页面图", document_id))

        page_count = max(len(page_texts), len(page_images))
        blocks = []
        sequence = 0
        for index in range(page_count):
            page = index + 1
            if index < len(page_images):
                sequence += 1
                blocks.append({
                    "type": "image", "path": self._relative(page_images[index]), "sha256": sha256_file(page_images[index]),
                    "mime_type": "image/png", "sequence": sequence, "page": page, "source_locator": f"page:{page}",
                })
            text = page_texts[index] if index < len(page_texts) else ""
            if text:
                sequence += 1
                blocks.append({"type": "text", "text": text, "sequence": sequence, "page": page, "source_locator": f"page:{page}"})
            else:
                issues.append(self._issue("warning", "pdf_page_text_missing", f"第 {page} 页没有可用文本，需 OCR 或人工核对", document_id))
        return blocks, issues

    def _parse_xlsx(self, source: Path, document_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        if not importlib.util.find_spec("openpyxl"):
            return [], [self._issue("error", "xlsx_dependency_missing", "缺少 openpyxl，无法读取 Excel", document_id)]
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            blocks = []
            for sequence, sheet in enumerate(workbook.worksheets, 1):
                rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
                blocks.append({"type": "table", "rows": rows, "sequence": sequence, "source_locator": f"sheet:{sheet.title}"})
            workbook.close()
            return blocks, []
        except Exception as error:
            return [], [self._issue("error", "xlsx_parse_failed", str(error), document_id)]

    def _segment_document(self, snapshot: dict[str, Any], document: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        blocks = document["content_blocks"]
        if not blocks:
            return [], []
        if any("page" in block for block in blocks):
            groups: list[list[dict[str, Any]]] = []
            for page in sorted({block.get("page", 1) for block in blocks}):
                groups.append([block for block in blocks if block.get("page", 1) == page])
        else:
            groups = []
            current: list[dict[str, Any]] = []
            for block in blocks:
                text = block.get("text", "") if block["type"] in {"text", "formula"} else ""
                if QUESTION_START.match(text) and current:
                    groups.append(current)
                    current = []
                current.append(block)
            if current:
                groups.append(current)
        assets = []
        issues = []
        for index, group in enumerate(groups, 1):
            asset = self._asset_from_blocks(snapshot, document, group, index, None)
            if asset["image_integrity"] == "missing":
                issue = self._issue("error", "declared_image_missing", "题干声明含图，但内容块中没有对应图片", document["id"], asset["id"])
                issues.append(issue)
                asset["issue_codes"].append(issue["code"])
            assets.append(asset)
        return assets, issues

    def _asset_from_blocks(
        self,
        snapshot: dict[str, Any],
        document: dict[str, Any],
        blocks: list[dict[str, Any]],
        ordinal: int,
        source_record_id: str | None,
        source_fields: dict[str, Any] | None = None,
        normalized_fields: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        texts = [str(block.get("text", "")) for block in blocks if block["type"] in {"text", "formula"}]
        for block in blocks:
            if block["type"] == "table":
                texts.append(json.dumps(block.get("rows", []), ensure_ascii=False))
        full_text = "\n".join(value for value in texts if value).strip()
        visual_hashes = [str(block.get("sha256") or block.get("value")) for block in blocks if block["type"] in {"image", "image_reference"}]
        fingerprint_source = normalized_text(full_text) or "|".join(visual_hashes)
        fingerprint = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
        match = QUESTION_START.match(full_text)
        question_no = match.group(1) if match else None
        asset_id = stable_id("question", snapshot["id"], document["id"], source_record_id or ordinal)
        units = []
        for unit_index, match in enumerate(UNIT_PATTERN.finditer(full_text), 1):
            label = match.group(1) or match.group(2)
            units.append({"id": f"{asset_id}-u{unit_index}", "label": label, "source_span": [match.start(), match.end()]})
        has_image = any(block["type"] == "image" for block in blocks)
        has_reference = any(block["type"] == "image_reference" for block in blocks)
        image_integrity = "preserved" if has_image else ("remote_reference_unmaterialized" if has_reference else ("missing" if IMAGE_HINT.search(full_text) else "no_visual_declared"))
        fields = dict(source_fields or {})
        fields.update({key: value for key, value in (normalized_fields or {}).items() if value not in (None, "", [])})
        units = attach_unit_tag_profiles(units, fields)
        return {
            "id": asset_id, "source_snapshot_id": snapshot["id"], "source_document_id": document["id"],
            "source_record_id": source_record_id, "source_name": document["source_name"], "ordinal": ordinal,
            "question_no": question_no, "title": (full_text[:80] or f"{document['source_name']} 第 {ordinal} 项").replace("\n", " "),
            "content_blocks": blocks, "units": units, "raw_text": full_text, "source_fields": source_fields or {},
            "normalized_fields": normalized_fields or {}, "fingerprint": fingerprint, "duplicate_group_id": None,
            "tag_profile": build_tag_profile(fields),
            "duplicate_count": 1, "image_integrity": image_integrity, "issue_codes": [], "status": "standardized",
            "parser_version": PARSER_VERSION,
        }

    @staticmethod
    def _text_blocks(text: str) -> list[dict[str, Any]]:
        parts = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        return [{"type": "text", "text": part, "sequence": index, "source_locator": f"text-block:{index}"} for index, part in enumerate(parts, 1)]

    @staticmethod
    def _issue(severity: str, code: str, message: str, document_id: str | None, question_asset_id: str | None = None) -> dict[str, Any]:
        return {
            "id": f"issue-{uuid4().hex[:12]}", "severity": severity, "code": code, "message": message,
            "source_document_id": document_id, "question_asset_id": question_asset_id,
        }

    @staticmethod
    def _mapped_value(fields: dict[str, Any], source_name: str | None) -> Any:
        return fields.get(source_name) if source_name else None

    @staticmethod
    def _plain_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, str):
            return value.strip()
        if isinstance(value, list):
            return "、".join(DocumentStandardizer._plain_value(item) for item in value)
        if isinstance(value, dict):
            return json.dumps(value, ensure_ascii=False)
        return str(value)

    def _relative(self, path: Path) -> str:
        return path.resolve().relative_to(self.output_root).as_posix()

    @staticmethod
    def _page_number(path: Path) -> int:
        match = re.search(r"-(\d+)\.png$", path.name)
        return int(match.group(1)) if match else 0
